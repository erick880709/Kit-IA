"""Indicadores del dashboard operativo y exportación de reportes (HU-E6-01/03).

Todos los indicadores se calculan desde los registros reales de la BD
(CA4, nada estático). Reportes anonimizados: sin identificadores directos.
"""

from __future__ import annotations

import csv
import io
import json

import pandas as pd
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.domain.catalogos import NIVELES_TRIaje
from app.domain.entities import EventoTriaje
from app.services import modelo_service

METAS_RNF001 = {"f1": 0.82, "precision": 0.85, "recall": 0.80, "auc_roc": 0.87}


def _metricas_modelo(session: Session) -> dict:
    activo = modelo_service.modelo_activo(session)
    if activo is None or not activo.metricas_json:
        return {}
    try:
        return json.loads(activo.metricas_json)
    except (json.JSONDecodeError, TypeError):
        return {}


def _tiempo_promedio_atencion(eventos: list[EventoTriaje]) -> float | None:
    duraciones = [
        (e.cierre - e.inicio).total_seconds() / 60
        for e in eventos
        if e.cierre is not None and e.inicio is not None
    ]
    return round(sum(duraciones) / len(duraciones), 1) if duraciones else None


def calcular_indicadores(session: Session) -> dict:
    """CA1: 7 indicadores + semáforo de metas + matriz de confusión."""
    eventos = list(session.scalars(select(EventoTriaje)).all())
    cerrados = [e for e in eventos if e.estado == "Cerrado"]

    nivel_profesional = [e.nivel_asignado_profesional for e in cerrados]
    distribucion = {
        n: nivel_profesional.count(n) for n in NIVELES_TRIaje
    }
    total_dist = sum(distribucion.values()) or 1

    concordancias = [e.concordancia for e in cerrados if e.concordancia is not None]
    concordancia_global = (
        round(sum(concordancias) / len(concordancias), 3) if concordancias else None
    )

    concordancia_por_nivel = {}
    for n in NIVELES_TRIaje:
        fila = [
            e for e in cerrados
            if e.nivel_asignado_profesional == n and e.concordancia is not None
        ]
        concordancia_por_nivel[n] = (
            round(sum(e.concordancia for e in fila) / len(fila), 3) if fila else None
        )

    confusion = pd.crosstab(
        pd.Series([e.nivel_sugerido_ia or "—" for e in cerrados], name="IA"),
        pd.Series([e.nivel_asignado_profesional or "—" for e in cerrados],
                 name="Profesional"),
    ).reindex(index=NIVELES_TRIaje, columns=NIVELES_TRIaje, fill_value=0)

    discrepancias = [
        {
            "fecha": e.cierre.strftime("%Y-%m-%d %H:%M") if e.cierre else "—",
            "nivel_ia": e.nivel_sugerido_ia or "—",
            "nivel_profesional": e.nivel_asignado_profesional or "—",
            "motivo": e.motivo_discrepancia or "—",
        }
        for e in cerrados
        if e.concordancia is False
    ]

    metricas = _metricas_modelo(session)
    macro = metricas.get("macro", {})
    desempeno = {
        "version": modelo_service.modelo_activo(session).version
        if modelo_service.modelo_activo(session) else "—",
        "precision": macro.get("precision"),
        "recall": macro.get("recall"),
        "f1": macro.get("f1"),
        "auc": metricas.get("auc_roc_ovr"),
    }

    # CA2: semáforo de metas RNF-001
    semaforo = {}
    for clave, meta in METAS_RNF001.items():
        clave_desempeno = "auc" if clave == "auc_roc" else clave
        valor = desempeno.get(clave_desempeno)
        if valor is None:
            semaforo[clave] = {"meta": meta, "valor": None, "estado": "sin_dato"}
        else:
            semaforo[clave] = {
                "meta": meta, "valor": round(float(valor), 4),
                "estado": "ok" if float(valor) >= meta else "alerta",
            }

    return {
        "n_eventos": len(eventos),
        "n_cerrados": len(cerrados),
        "distribucion": {n: v / total_dist for n, v in distribucion.items()},
        "tiempo_promedio_atencion_min": _tiempo_promedio_atencion(eventos),
        "desempeno_ia": desempeno,
        "concordancia_global": concordancia_global,
        "concordancia_por_nivel": concordancia_por_nivel,
        "matriz_confusion": confusion,
        "discrepancias": discrepancias,
        "semaforo": semaforo,
    }


def _hoja_exportacion(indicadores: dict) -> list[dict]:
    sem = indicadores["semaforo"]
    base = [
        {"indicador": "Volumen de eventos", "valor": indicadores["n_eventos"]},
        {"indicador": "Eventos cerrados", "valor": indicadores["n_cerrados"]},
        {"indicador": "Tiempo promedio de atención (min)",
         "valor": indicadores["tiempo_promedio_atencion_min"]},
        {"indicador": "Concordancia global IA vs profesional",
         "valor": indicadores["concordancia_global"]},
        {"indicador": "F1 (macro)", "valor": sem["f1"]["valor"],
         "meta": sem["f1"]["meta"], "estado": sem["f1"]["estado"]},
        {"indicador": "Precisión (macro)", "valor": sem["precision"]["valor"],
         "meta": sem["precision"]["meta"], "estado": sem["precision"]["estado"]},
        {"indicador": "Recall (macro)", "valor": sem["recall"]["valor"],
         "meta": sem["recall"]["meta"], "estado": sem["recall"]["estado"]},
        {"indicador": "AUC-ROC (OVR)", "valor": sem["auc_roc"]["valor"],
         "meta": sem["auc_roc"]["meta"], "estado": sem["auc_roc"]["estado"]},
        {"indicador": "Distribución I", "valor": indicadores["distribucion"].get("I")},
        {"indicador": "Distribución II", "valor": indicadores["distribucion"].get("II")},
        {"indicador": "Distribución III", "valor": indicadores["distribucion"].get("III")},
        {"indicador": "Distribución IV", "valor": indicadores["distribucion"].get("IV")},
        {"indicador": "Distribución V", "valor": indicadores["distribucion"].get("V")},
    ]
    for fila in base:
        fila.setdefault("meta", None)
        fila.setdefault("estado", None)
    return base


def _excel_con_graficos(base: list[dict], indicadores: dict, tendencia: list[dict] | None) -> bytes:
    """Excel con hojas 'dashboard' + 'graficos' (gráficos nativos de Excel)."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference

    wb = Workbook()
    ws = wb.active
    ws.title = "dashboard"
    ws.append(list(base[0].keys()))
    for fila in base:
        ws.append([fila.get(c) for c in base[0]])

    g = wb.create_sheet("graficos")
    # Distribución por nivel
    g["A1"] = "Nivel"
    g["B1"] = "Proporcion"
    dist = indicadores["distribucion"]
    for i, nivel in enumerate(NIVELES_TRIaje):
        g[f"A{i + 2}"] = nivel
        g[f"B{i + 2}"] = round(float(dist.get(nivel, 0)), 4)
    bar = BarChart()
    bar.title = "Distribucion de triaje por nivel"
    bar.y_axis.title = "Proporcion"
    bar.add_data(Reference(g, min_col=2, min_row=1, max_row=1 + len(NIVELES_TRIaje)),
                 titles_from_data=True)
    bar.set_categories(Reference(g, min_col=1, min_row=2, max_row=1 + len(NIVELES_TRIaje)))
    g.add_chart(bar, "D2")

    # Concordancia por nivel
    inicio = 9
    g[f"A{inicio}"] = "Nivel"
    g[f"B{inicio}"] = "Concordancia"
    conc = indicadores["concordancia_por_nivel"]
    for i, nivel in enumerate(NIVELES_TRIaje):
        g[f"A{inicio + 1 + i}"] = nivel
        g[f"B{inicio + 1 + i}"] = (
            round(float(conc[nivel]), 4) if conc.get(nivel) is not None else 0.0
        )
    bar2 = BarChart()
    bar2.title = "Concordancia IA vs profesional por nivel"
    bar2.add_data(
        Reference(g, min_col=2, min_row=inicio, max_row=inicio + len(NIVELES_TRIaje)),
        titles_from_data=True,
    )
    bar2.set_categories(
        Reference(g, min_col=1, min_row=inicio + 1, max_row=inicio + len(NIVELES_TRIaje))
    )
    g.add_chart(bar2, "D16")

    # Tendencia diaria (si hay serie)
    if tendencia:
        ini2 = 18
        g[f"A{ini2}"] = "Fecha"
        g[f"B{ini2}"] = "Eventos"
        for i, fila in enumerate(tendencia):
            g[f"A{ini2 + 1 + i}"] = str(fila["fecha"])
            g[f"B{ini2 + 1 + i}"] = int(fila["n"])
        linea = LineChart()
        linea.title = "Tendencia diaria de eventos"
        linea.add_data(
            Reference(g, min_col=2, min_row=ini2, max_row=ini2 + len(tendencia)),
            titles_from_data=True,
        )
        linea.set_categories(
            Reference(g, min_col=1, min_row=ini2 + 1, max_row=ini2 + len(tendencia))
        )
        g.add_chart(linea, "D30")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _barras_pdf(pdf, titulo: str, datos: dict, x: float, y: float, ancho: float) -> float:
    """Barras horizontales simples dibujadas en el PDF. Devuelve la nueva y."""
    pdf.setFont("Helvetica-Bold", 11)
    pdf.setFillColor("#164E63")
    pdf.drawString(x, y, titulo)
    y -= 0.8 * cm
    valores = {k: float(v or 0.0) for k, v in datos.items()}
    maximo = max(valores.values()) or 1.0
    for clave, valor in valores.items():
        pdf.setFont("Helvetica", 9)
        pdf.setFillColor("#164E63")
        pdf.drawString(x, y + 0.12 * cm, clave)
        ancho_barra = (valor / maximo) * (ancho * 0.72)
        pdf.setFillColor("#0891B2")
        pdf.roundRect(x + 1.2 * cm, y, ancho_barra, 0.5 * cm, 2, stroke=0, fill=1)
        pdf.setFillColor("#164E63")
        pdf.drawString(x + 1.4 * cm + ancho_barra, y + 0.12 * cm, f"{valor:.3f}")
        y -= 0.7 * cm
    return y - 0.3 * cm


def _tendencia_pdf(
    pdf, tendencia: list[dict], x: float, y: float, ancho: float, alto: float
) -> float:
    """Línea de tendencia diaria dibujada en el PDF. Devuelve la nueva y."""
    pdf.setFont("Helvetica-Bold", 11)
    pdf.setFillColor("#164E63")
    pdf.drawString(x, y, "Tendencia diaria de eventos (14 dias)")
    y -= 0.5 * cm
    if not tendencia:
        return y - 0.5 * cm
    maximo = max(int(f["n"]) for f in tendencia) or 1
    n = len(tendencia)
    paso = ancho / max(n - 1, 1)
    puntos = [
        (x + i * paso, y + (int(f["n"]) / maximo) * alto)
        for i, f in enumerate(tendencia)
    ]
    pdf.setStrokeColor("#0891B2")
    pdf.setLineWidth(1.4)
    p = pdf.beginPath()
    p.moveTo(*puntos[0])
    for px, py in puntos[1:]:
        p.lineTo(px, py)
    pdf.drawPath(p, stroke=1, fill=0)
    pdf.setStrokeColor("#94A3B8")
    pdf.setLineWidth(0.8)
    pdf.line(x, y, x + ancho, y)
    pdf.setFont("Helvetica", 8)
    pdf.setFillColor("#164E63")
    pdf.drawString(x, y - 0.4 * cm, str(tendencia[0]["fecha"]))
    pdf.drawRightString(x + ancho, y - 0.4 * cm, str(tendencia[-1]["fecha"]))
    return y - 0.9 * cm


def exportar_reporte(
    indicadores: dict, *, formato: str, tendencia: list[dict] | None = None,
) -> tuple[bytes, str]:
    """HU-E6-03 CA1/CA2: exportación CSV/Excel/PDF anonimizada CON gráficos.

    Excel incluye gráficos nativos (barras y línea) en la hoja «graficos»;
    el PDF incluye los gráficos dibujados además de la tabla de indicadores.
    """
    base = _hoja_exportacion(indicadores)
    if formato == "csv":
        buffer = io.StringIO()
        campos = list(base[0].keys())
        escritor = csv.DictWriter(buffer, fieldnames=campos)
        escritor.writeheader()
        escritor.writerows(base)
        return buffer.getvalue().encode("utf-8-sig"), "reporte_dashboard.csv"
    if formato == "excel":
        return _excel_con_graficos(base, indicadores, tendencia), "reporte_dashboard.xlsx"
    if formato == "pdf":
        buffer = io.BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=A4, pageCompression=0)
        w, h = A4
        pdf.setFillColor("#0891B2")
        pdf.rect(0, h - 2.5 * cm, w, 2.5 * cm, stroke=0, fill=1)
        pdf.setFillColor("#FFFFFF")
        pdf.setFont("Helvetica-Bold", 16)
        pdf.drawString(2 * cm, h - 1.6 * cm, "TriajeIA — Reporte operativo")
        pdf.setFillColor("#164E63")
        y = h - 3.6 * cm
        pdf.setFont("Helvetica", 10)
        for fila in base:
            pdf.drawString(
                2 * cm, y,
                f"{fila['indicador']}: {fila['valor']}"
                + (f" · meta {fila['meta']} · {fila['estado']}"
                   if fila.get("meta") is not None else ""),
            )
            y -= 0.6 * cm
            if y < 2 * cm:
                pdf.showPage()
                y = h - 2.5 * cm
        # Gráfico 1: distribución por nivel
        if y < 7 * cm:
            pdf.showPage()
            y = h - 2.5 * cm
        y = _barras_pdf(
            pdf, "Distribucion de triaje por nivel", indicadores["distribucion"],
            2 * cm, y - 0.6 * cm, w - 4 * cm,
        )
        # Gráfico 2: concordancia por nivel
        if y < 6.5 * cm:
            pdf.showPage()
            y = h - 2.5 * cm
        y = _barras_pdf(
            pdf, "Concordancia IA vs profesional por nivel",
            indicadores["concordancia_por_nivel"], 2 * cm, y - 0.3 * cm, w - 4 * cm,
        )
        # Gráfico 3: tendencia diaria
        if y < 8 * cm:
            pdf.showPage()
            y = h - 2.5 * cm
        y = _tendencia_pdf(pdf, tendencia or [], 2 * cm, y - 0.3 * cm, w - 4 * cm, 4 * cm)
        pdf.drawString(
            2 * cm, y - 0.4 * cm,
            "Reporte anonimizado — sin identificadores directos de pacientes.",
        )
        pdf.showPage()
        pdf.save()
        return buffer.getvalue(), "reporte_dashboard.pdf"
    raise ValueError(f"Formato no soportado: {formato}")


def conteo_por_dia(session: Session, dias: int = 14) -> list[dict]:
    """Serie diaria de eventos (para tendencia en el dashboard).

    Usa el helper único de "ahora" en UTC naive (sin mezcla de zonas).
    """
    from datetime import timedelta

    from app.services.audit_service import ahora_utc_naive

    desde = ahora_utc_naive() - timedelta(days=dias)
    filas = session.execute(
        select(func.date(EventoTriaje.inicio), func.count(EventoTriaje.id))
        .where(EventoTriaje.inicio >= desde)
        .group_by(func.date(EventoTriaje.inicio))
        .order_by(func.date(EventoTriaje.inicio))
    ).all()
    return [{"fecha": str(f), "n": int(n)} for f, n in filas]
